"""Row-list serializers: CSV, JSON, Markdown, and multi-sheet Excel."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

from p6_mcp.exceptions import OptionalDependencyError
from p6_mcp.services.query.serialize import iso


def _columns(rows: list[dict[str, Any]]) -> list[str]:
    """Union of keys in first-seen order (stable across rows)."""
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols


def _flat(value: Any) -> Any:
    """Flatten nested structures for tabular output."""
    v = iso(value)
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, default=str)
    return v


def to_csv(rows: list[dict[str, Any]]) -> str:
    """Rows as CSV text (RFC 4180, CRLF suppressed to \\n)."""
    if not rows:
        return ""
    cols = _columns(rows)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=cols, lineterminator="\n", extrasaction="ignore")
    writer.writeheader()
    for r in rows:
        writer.writerow({c: _flat(r.get(c)) for c in cols})
    return buf.getvalue()


def to_json_text(data: Any, indent: int = 2) -> str:
    """Any structure as pretty JSON with ISO dates."""
    return json.dumps(iso(data), indent=indent, ensure_ascii=False, default=str)


def to_markdown(rows: list[dict[str, Any]], max_rows: int | None = None) -> str:
    """Rows as a GitHub-flavored Markdown table."""
    if not rows:
        return "_(no rows)_"
    cols = _columns(rows)
    shown = rows[:max_rows] if max_rows else rows
    lines = [
        "| " + " | ".join(cols) + " |",
        "| " + " | ".join("---" for _ in cols) + " |",
    ]
    for r in shown:
        cells = []
        for c in cols:
            text = "" if r.get(c) is None else str(_flat(r.get(c)))
            cells.append(text.replace("|", "\\|").replace("\n", " "))
        lines.append("| " + " | ".join(cells) + " |")
    if max_rows and len(rows) > max_rows:
        lines.append(f"\n_{len(rows) - max_rows} more rows not shown._")
    return "\n".join(lines)


def to_excel(sheets: dict[str, list[dict[str, Any]]], output_path: Path) -> Path:
    """Write a multi-sheet .xlsx workbook (requires the ``excel`` extra)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover - exercised via error path
        raise OptionalDependencyError(
            "Excel export requires openpyxl",
            hint="Install with: pip install 'p6-mcp[excel]'",
        ) from exc
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name[:31] or "Sheet")
        if not rows:
            ws.append(["(no rows)"])
            continue
        cols = _columns(rows)
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([_flat(r.get(c)) for c in cols])
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, start=1):
            width = max(len(c), *(len(str(_flat(r.get(c)) or "")) for r in rows[:200]))
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(
                max(width + 2, 10), 55
            )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path.resolve()
